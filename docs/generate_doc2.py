from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_margins(section):
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for sec in doc.sections:
    set_margins(sec)

def add_heading(text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.font.all_caps = True
    elif level == 2:
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        run.font.size = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run.font.name = 'Times New Roman'
    return p

def add_body(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.name = 'Times New Roman'
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'Times New Roman'
    doc.add_paragraph()
    return table

# Title
for i in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("SHOPSMART")
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x0E, 0x4A, 0x8A)
run.font.name = 'Times New Roman'
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("AI-Enabled E-Commerce Price Comparison Platform")
run.font.size = Pt(14)
run.font.name = 'Times New Roman'
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("A PROJECT REPORT\nSubmitted in partial fulfillment of\nBachelor of Computer Applications (BCA)\nFinal Year Project 2025-26")
run.font.size = Pt(12)
run.font.name = 'Times New Roman'
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Submitted By:\n[Student Name]  [Roll No]\nUnder the Guidance of\n[Guide Name]\n\n[College Name]\nAffiliated to [University Name]\nAcademic Year 2025-26")
run.font.size = Pt(11)
run.font.name = 'Times New Roman'
doc.add_page_break()

add_heading("CERTIFICATE OF AUTHENTICITY", level=1)
add_body("This is to certify that the project entitled SHOPSMART – AI-Enabled E-Commerce Price Comparison Platform is a bonafide work carried out by the student of BCA Final Year in partial fulfillment of the requirements for the award of Bachelor of Computer Applications. The project has been carried out under our supervision and guidance and no part of this report has been submitted elsewhere for any other degree or diploma.")
add_body("The project demonstrates original work in full-stack development using React, Django REST Framework and PostgreSQL with AI-enabled features for price comparison across Amazon, Flipkart, Croma and other platforms.")
p = doc.add_paragraph()
p.add_run("\n\nDate: 02-09-2026\nPlace: [College City]\n\n\n______________________\t\t______________________\nProject Guide\t\t\t\tHead of Department\n[Guide Name]\t\t\t\t[HoD Name]\n").font.size = Pt(11)
doc.add_page_break()

add_heading("ACKNOWLEDGMENTS", level=1)
add_body("We express our sincere gratitude to our project guide and the Department of Computer Applications for their continuous support, encouragement and valuable guidance throughout the development of ShopSmart. We are thankful to the faculty members, our parents and peers for their motivation and to the open-source community (Django, React, PostgreSQL) for excellent documentation. We also thank the college administration for providing necessary infrastructure.")
add_body("This project has enhanced our practical knowledge of full-stack architecture, database design, REST APIs, authentication (JWT), and AI integration for e-commerce analytics.")
doc.add_page_break()

add_heading("ABSTRACT", level=1)
add_body("ShopSmart is an AI-enabled e-commerce price comparison platform that helps users find the best deal by aggregating prices, ratings, discounts, coupons, delivery and availability across eight major Indian platforms: Amazon, Flipkart, AJIO, Myntra, Meesho, Croma, Reliance Digital and Tata CLiQ. The system is built as a professional full-stack application with React + Vite frontend, Django + Django REST Framework backend and PostgreSQL database, secured with JWT authentication.")
add_body("Key features include product search by name/brand/category/platform, category filtering, product cards with platform badges, price comparison table with cheapest highlight, product details with specifications, real PriceHistory stored in PostgreSQL and visualized as a 30-day Recharts graph, persistent wishlist (guest + authenticated), price-drop alerts, price history management via scrapers (Amazon/Flipkart with fallback jitter), user profile/settings, admin module (IsAdminUser), and AI verdict/sentiment. Demo data comprises 240 products across 16 categories seeded via Django management command with offers (815 rows) and history (5705 rows).")
add_body("Keywords: Price Comparison, E-Commerce, Django REST, React, PostgreSQL, JWT, Web Scraping, Price History, Recharts, AI Recommendation")
doc.add_page_break()

add_heading("TABLE OF CONTENTS", level=1)
toc_items = ["Preliminary Pages - i", "Certificate of Authenticity - ii", "Acknowledgments - iii", "Abstract - iv", "List of Figures - v", "List of Tables - vi", "Chapter 1: Introduction - 1", "Chapter 2: System Analysis - 4", "Chapter 3: Software Requirement Specifications - 6", "Chapter 4: System Design - 8", "Chapter 5: Implementation & Coding - 12", "Chapter 6: System Testing - 16", "Chapter 7: Conclusion & Future Enhancements - 18", "Appendices & References - 19"]
for item in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(item)
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'

add_heading("CHAPTER 1: INTRODUCTION", level=1)
add_heading("1.1 Project Overview", level=2)
add_body("ShopSmart is a BCA final-year full-stack project that aggregates product information across multiple e-commerce platforms into a single price-comparison interface. Users can search products, filter by category, view price comparison tables, track 30-day price history graphs, manage wishlist and set price-drop alerts. The application is designed to be functional, professional, demo-friendly and maintainable for academic evaluation.")
add_heading("1.2 Motivation & Domain Description", level=2)
add_body("Indian e-commerce shoppers routinely compare prices manually across Amazon, Flipkart, Myntra, etc., which is time-consuming and error-prone. The price-comparison domain requires aggregating heterogeneous product data (price, rating, delivery, availability) and presenting a unified view with historical trends. Motivation is to reduce shopping effort, increase transparency and apply AI for deal recommendation.")
add_heading("1.3 Problem Statement", level=2)
add_body("Existing manual comparison lacks centralized view, historical price insight, and personalized alerts. Users cannot easily identify the cheapest valid offer with delivery/coupon context. Previous ShopSmart prototype had only 2 demo products, fake price graphs, minimal ProductOffer fields, no JWT auth, no pagination and Vite EPERM file-lock issues on Windows.")
add_heading("1.4 Objectives", level=2)
for obj in ["Provide unified price comparison across 8 platforms", "Implement real PriceHistory in PostgreSQL with 30-day Recharts graph", "Support search, category filtering, sorting, pagination (24/page)", "Enable persistent wishlist (guest + JWT) and price-drop alerts", "Deliver modern responsive UI", "Ensure clean Django architecture with safe migrations and JWT auth"]:
    add_bullet(obj)
add_heading("1.5 Scope", level=2)
add_body("Scope covers 240 demo products across 16 categories seeded via Django command, REST APIs for products/offers/history/categories/wishlist/auth/alerts, scrapers for Amazon/Flipkart with fallback jitter, and admin module for product/user management. Out of scope are payment gateway, real-time third-party APIs, and email/SMS notifications.")

add_heading("CHAPTER 2: SYSTEM ANALYSIS", level=1)
add_heading("2.1 Existing System Study", level=2)
add_body("Existing price-comparison sites (PriceDekho, MySmartPrice) aggregate via affiliate APIs. Our earlier ShopSmart v0 had Django backend with SQLite, minimal offers, and React frontend with hardcoded products and random-value graphs.")
add_heading("2.2 Limitations of Existing System", level=2)
for lim in ["Manual tab-switching, no 30-day trend", "No persistent wishlist/alerts", "Only 2 products, not 16 categories", "PriceHistory schema mismatch", "Vite EPERM rename error on Windows", "No admin controls"]:
    add_bullet(lim)
add_heading("2.3 Proposed System Features", level=2)
add_body("Proposed ShopSmart v1 addresses all limitations with full-stack flow: PostgreSQL to Django Models to DRF to React to ShopSmart UI. Core modules: Home, Navbar, Search, Categories, Product Cards, Product Details/Compare, Price History Graph, Wishlist, Alerts, Auth, Profile, Admin, AI verdict.")
add_heading("2.4 Feasibility Study", level=2)
add_heading("2.4.1 Technical Feasibility", level=3)
add_body("Stack: React 19 + Vite 8, Django 6.0.7 + DRF 3.17 + SimpleJWT 5.5, PostgreSQL 15+psycopg2, Recharts 3.10, BeautifulSoup4 4.15. All open-source, Windows/Linux compatible.")
add_heading("2.4.2 Operational Feasibility", level=3)
add_body("Operational: Single runserver 127.0.0.1:8000 + npm run dev at 5173. Demo data via seed_products --count 240. Admin via /admin and /admin frontend.")
add_heading("2.4.3 Economic Feasibility", level=3)
add_body("Economic: Zero licensing cost. Hosting on free tiers minimal. No paid APIs; scrapers with fallback jitter avoid cost.")

add_heading("CHAPTER 3: SOFTWARE REQUIREMENT SPECIFICATIONS (SRS)", level=1)
add_heading("3.1 Functional Requirements", level=2)
add_table(["ID", "Requirement", "Priority"], [["FR1", "User registration/login with JWT (min 8 chars, Eye toggle, forgot-password)", "High"], ["FR2", "Search by name/brand/category/platform", "High"], ["FR3", "Category filter (16 values) via GET /api/categories/", "High"], ["FR4", "Product listing paginated (24/page)", "High"], ["FR5", "Product details + comparison table with cheapest highlight", "High"], ["FR6", "PriceHistory 30-day per-store graph (5705 rows)", "High"], ["FR7", "Wishlist add/remove/persistent (guest_id + JWT)", "High"], ["FR8", "Price drop alerts (target_price, is_triggered)", "Medium"], ["FR9", "Profile edit, change password, admin stats/delete", "Medium"], ["FR10", "Buy via store search URL soft-gate if not logged in", "Medium"]])
add_heading("3.2 Non-Functional Requirements", level=2)
for nfr in ["Performance: API <200ms for 24 items", "Security: JWT 60m/7d, CORS, password min 8", "Usability: Responsive grid, dark/light, hover, loading/empty/error states", "Maintainability: Modular Django apps, safe migrations", "Portability: Windows Vite EPERM handled"]:
    add_bullet(nfr)
add_heading("3.3 Hardware Requirements", level=2)
add_table(["Component", "Specification"], [["CPU", "Intel i3 / Ryzen 3+"], ["RAM", "8 GB"], ["Storage", "5 GB free"], ["Network", "Broadband"]])
add_heading("3.4 Software Stack Configuration", level=2)
add_table(["Layer", "Tech & Version", "Purpose"], [["Frontend", "React 19.2, Vite 8.2, recharts 3.10, axios 1.20", "UI, routing, graphs"], ["Backend", "Django 6.0.7, DRF 3.17, SimpleJWT 5.5", "REST, auth, scraping"], ["Database", "PostgreSQL 15, psycopg2", "Persistent storage"]])

add_heading("CHAPTER 4: SYSTEM DESIGN", level=1)
add_heading("4.1 High-Level Architecture", level=2)
add_body("Architecture: PostgreSQL to Django Models (Product, ProductOffer, PriceHistory, Wishlist, PriceAlert, User) to DRF Serializers/ViewSets to React (axios, AuthContext) to ShopSmart UI. Scrapers to update_prices command to PriceHistory.")
add_heading("4.2 Data Flow Diagrams", level=2)
add_heading("4.2.1 DFD Level 0 (Context)", level=3)
add_body("External entities: User, Admin, E-commerce Stores. Process: ShopSmart System. Flows: User Search/Compare/Wishlist/Alerts; Admin Manage; Stores Offers/Price via Scrapers.")
add_heading("4.2.2 DFD Level 1", level=3)
add_body("Processes: 1. Auth, 2. Product Management, 3. Price Aggregation, 4. User Features, 5. Admin. Datastores: User, Product, Offer, History, Wishlist, Alert.")
add_heading("4.2.3 DFD Level 2 (Product Detail)", level=3)
add_body("User requests /product/:id to DRF retrieve + history action + compare action (cheapest calc).")
add_heading("4.3 ER Diagrams", level=2)
add_body("Entities: Product 1-M ProductOffer, Product 1-M PriceHistory, Product 1-M Wishlist, User 1-M Wishlist, User 1-M PriceAlert, Product 1-M PriceAlert. ProductOffer unique(product, store_name).")
add_heading("4.4 Database Schema", level=2)
add_table(["Table", "Columns", "Constraints"], [["auth_user", "id PK, username, email, password, is_staff, date_joined", "Unique username/email"], ["products_product", "id PK, title, category, brand, image_url, description, specifications JSON, created_at, updated_at", "ordering -created_at"], ["products_productoffer", "id PK, product_id FK, store_name, current_price, original_price, discount, rating, reviews_count, coupon_code, product_url, in_stock, delivery_days", "unique(product,store_name)"], ["products_pricehistory", "id PK, product_id FK, store_name, price, recorded_at", "ordering -recorded_at"], ["products_wishlist", "id PK, user_id FK nullable, guest_id, product_id FK, added_at", "ordering -added_at"], ["users_pricealert", "id PK, user_id FK, product_id FK, target_price, active, created_at", "ordering -created_at"]])
add_heading("4.5 UI Mockups", level=2)
add_body("Home: hero Compare Prices Across Stores + search + chips + grid (cards with brand tag, image, lowest price, savings pill, 4 offers, Graph/AI/Details) + top+bottom pagination; Detail: image 360px + specs + best price + coupon + Buy soft-gate modal + alert box + comparison table BEST DEAL green + 30-day multi-line graph.")

add_heading("CHAPTER 5: IMPLEMENTATION & CODING", level=1)
add_heading("5.1 Module Breakdown", level=2)
add_table(["Module", "Files", "Responsibility"], [["Config", "config/settings.py, urls.py", "Installed apps, REST_FRAMEWORK+JWT, CORS, DB"], ["Products", "models.py, serializers.py, views.py, scrapers", "Product/Offer/History/Wishlist, history/compare actions"], ["Users", "models.py, serializers.py, views.py", "User+PriceAlert, JWT"], ["Frontend Core", "App.jsx, api/axios.js, context/AuthContext", "Router, JWT interceptor, guest_id"], ["Pages", "Home, ProductDetail, Wishlist, Alerts, Profile, Admin, ForgotPassword", "Features"]])
add_heading("5.2 Core Algorithms", level=2)
add_heading("5.2.1 Price Aggregation & Fallback", level=3)
add_body("Base scraper get_fallback_price generates +-1-2 percent jitter when CAPTCHA triggers. AmazonScraper parses a-price-whole, Flipkart parses _30jeq3. update_prices command updates current_price and creates PriceHistory. ProductOffer.save calculates discount_percent.")
add_heading("5.2.2 Cheapest & Savings", level=3)
add_body("ProductSerializer.get_lowest_price filters in_stock offers ordered by current_price. Frontend Detail computes cheapest = min(current_price where in_stock), highlights BEST DEAL.")
add_heading("5.2.3 Price History Grouping", level=3)
add_body("ProductViewSet.history filters PriceHistory by recorded_at asc. Frontend groups by toLocaleDateString into date, Amazon, Flipkart for Recharts Line per store with storeColors. YAxis domain auto prevents negative for low-price grocery.")
add_heading("5.3 Primary Code Snippets", level=2)
add_body("Snippet 1 - JWT Interceptor (api/axios.js): Skip X-Guest-ID for login/register to avoid CORS preflight, attach Bearer, handle 401 refresh. Snippet 2 - Wishlist Toggle (ProductCard.jsx): if wishlisted DELETE else POST with guest_id dispatch wishlist changed. Snippet 3 - Forgot Password (users/views.py): POST forgot-password verifies User.get(username,email) then set_password. Snippet 4 - Seed (seed_products.py): IMAGE_MAP per category for real product images, STORE_POOL per category, _history creates 7 points 30 to 0 days with recorded_at explicit.")

add_heading("CHAPTER 6: SYSTEM TESTING", level=1)
add_heading("6.1 Test Strategy", level=2)
add_body("Unit testing for serializers/models, integration via shell and APIRequestFactory, manual browser testing, smoke via Invoke-RestMethod. Focus on wishlist IntegrityError fix, PriceHistory auto_now_add fix, pagination.")
add_heading("6.2 Test Plan", level=2)
add_body("Test Plan: Backend check + showmigrations, GET products 240, GET categories 16, POST register min 8, POST login token, POST wishlist guest+auth, GET history 21 rows spanning 30 days, Buy search URL valid, POST forgot-password, Frontend build 0 errors.")
add_heading("6.3 Black-Box Test Cases Matrix", level=2)
add_table(["Test ID", "Input", "Expected Output", "Actual Result", "Status"], [["TC01", "GET /api/products/?page=1", "200 count 240 results 24", "200 count 240", "Pass"], ["TC02", "POST register short password", "400 min 8", "400 Ensure at least 8", "Pass"], ["TC03", "POST register Pktest01", "201 with tokens", "201 tokens", "Pass"], ["TC04", "POST login test_api_user", "200 access", "200 access 231 chars", "Pass"], ["TC05", "POST wishlist 508", "201", "201 id 14", "Pass"], ["TC06", "GET history 507", "21 rows oldest 2026-08-03", "21 oldest 2026-08-03", "Pass"], ["TC07", "POST forgot-password", "200 reset", "200 Password reset", "Pass"], ["TC08", "Buy guest", "Login modal", "Modal shows", "Pass"], ["TC09", "Admin non-staff", "Access denied + login form", "Pink error + form", "Pass"], ["TC10", "npm run build", "0 errors", "0 errors 712kB", "Pass"]])

add_heading("CHAPTER 7: CONCLUSION & FUTURE ENHANCEMENTS", level=1)
add_heading("7.1 Project Summary", level=2)
add_body("ShopSmart successfully delivers a demo-ready price-comparison platform with 240 products, real PriceHistory, JWT auth, wishlist/alerts, profile/admin, and modern UI. All previous issues are resolved. Documentation follows institutional 16/14/12 font, 1.5 spacing, 1.5/1.0 margins.")
add_heading("7.2 Constraints & Limitations", level=2)
for lim in ["Demo images are Unsplash/Amazon representative, not live scraped per product", "No email service for alerts (in-app only)", "No payment, affiliate checkout is search-URL soft-gate", "History is seeded, not live cron"]:
    add_bullet(lim)
add_heading("7.3 Future Enhancements", level=2)
for fut in ["Integrate Amazon Product Advertising API for live image/price", "AI Shopping Assistant chatbot (Gemini) + Deal Recommendation score", "Email/SMS price-drop notifier via Celery beat", "Admin CSV bulk upload + image upload", "PWA, caching, Elasticsearch", "User reviews/ratings with sentiment AI"]:
    add_bullet(fut)

add_heading("APPENDICES & REFERENCES", level=1)
add_heading("Bibliography (IEEE Format)", level=2)
for ref in ["[1] Django Documentation, https://docs.djangoproject.com/en/6.0/", "[2] Django REST Framework, https://www.django-rest-framework.org/", "[3] React Documentation, https://react.dev/", "[4] PostgreSQL Docs, https://www.postgresql.org/docs/", "[5] Unsplash Images, https://unsplash.com", "[6] Amazon Product Images (demo under fair use)"]:
    add_body(ref)
add_heading("User Manual", level=2)
add_body("1. Backend: cd backend venv\\Scripts\\python.exe manage.py runserver 127.0.0.1:8000 (DB: ShopSmart_db / postgres/Tony). 2. Frontend: cd frontend npm install npm run dev -> http://localhost:5173. 3. Demo users: Pktest01 / Adminttest@01 (staff), admin / Admin@123. 4. Forgot password: Login -> Forgot password -> username+email+new (min8). 5. Buy: ProductDetail -> Buy (guest shows login modal). 6. Admin: /admin -> login as staff -> stats + delete.")
add_heading("Abbreviations", level=2)
add_table(["Abbr", "Full Form"], [["BCA","Bachelor of Computer Applications"], ["DRF","Django REST Framework"], ["JWT","JSON Web Token"], ["DFD","Data Flow Diagram"], ["ER","Entity Relationship"], ["SRS","Software Requirement Specification"]])

output_path = r"D:\Final Year Project\ShopSmart\docs\ShopSmart_BCA_Documentation.docx"
doc.save(output_path)
print(f"Saved to {output_path} paragraphs {len(doc.paragraphs)} tables {len(doc.tables)}")
