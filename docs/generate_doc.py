from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import datetime

def set_margins(section):
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def style_paragraph(p, font_size=12, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, caps=False, centered=False):
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(2)
    for run in p.runs:
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.name = 'Times New Roman'
        if color:
            run.font.color.rgb = color
        if caps:
            run.font.all_caps = True

doc = Document()

# Default font
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

sections = doc.sections
for sec in sections:
    set_margins(sec)

# --- Helper to add chapter heading ---
def add_chapter_heading(text, level=1):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(16)
        run.font.all_caps = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        run.font.size = Pt(14)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        run.font.size = Pt(12)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.space_before = Pt(12)
    for r in p.runs:
        r.font.name = 'Times New Roman'
    return p

def add_body(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = 'Times New Roman'
    p.paragraph_format.left_indent = Inches(0.25 + level*0.25)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Light Grid Accent 1'
    # Header
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.name = 'Times New Roman'
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'Times New Roman'
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

# ==================== TITLE PAGE ====================
for i in range(3):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("SHOPSMART")
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x0E, 0x4A, 0x8A)
run.font.name = 'Times New Roman'
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("AI-Enabled E-Commerce Price Comparison Platform")
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
run.font.name = 'Times New Roman'

doc.add_paragraph().add_run("").add_break()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("A PROJECT REPORT\nSubmitted in partial fulfillment of\nBachelor of Computer Applications (BCA)\nFinal Year Project 2025-26")
run.font.size = Pt(12)
run.font.name = 'Times New Roman'
p.paragraph_format.line_spacing = 1.5

doc.add_paragraph().add_run("").add_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("Submitted By:\n[Student Name]  [Roll No]\nUnder the Guidance of\n[Guide Name]\n\n[College Name]\nAffiliated to [University Name]\nAcademic Year 2025-26")
run.font.size = Pt(11)
run.font.name = 'Times New Roman'

# Footer page number (roman for prelims)
sectPr = doc.sections[0]._sectPr
pgNumType = OxmlElement('w:pgNumType')
pgNumType.set(qn('w:fmt'), 'lowerRoman')
pgNumType.set(qn('w:start'), '1')
sectPr.append(pgNumType)

doc.add_page_break()

# ==================== CERTIFICATE ====================
add_chapter_heading("CERTIFICATE OF AUTHENTICITY", level=1)
add_body("This is to certify that the project entitled “SHOPSMART – AI-Enabled E-Commerce Price Comparison Platform” is a bonafide work carried out by the student of BCA Final Year in partial fulfillment of the requirements for the award of Bachelor of Computer Applications. The project has been carried out under our supervision and guidance and no part of this report has been submitted elsewhere for any other degree or diploma.")
add_body("The project demonstrates original work in full-stack development using React, Django REST Framework and PostgreSQL with AI-enabled features for price comparison across Amazon, Flipkart, Croma and other platforms.")
p = doc.add_paragraph()
p.add_run("\n\nDate: 02-09-2026\nPlace: [College City]\n\n\n______________________\t\t______________________\nProject Guide\t\t\t\tHead of Department\n[Guide Name]\t\t\t\t[HoD Name]\n").font.size = Pt(11)

doc.add_page_break()

# ==================== ACKNOWLEDGMENT ====================
add_chapter_heading("ACKNOWLEDGMENTS", level=1)
add_body("We express our sincere gratitude to our project guide and the Department of Computer Applications for their continuous support, encouragement and valuable guidance throughout the development of ShopSmart. We are thankful to the faculty members, our parents and peers for their motivation and to the open-source community (Django, React, PostgreSQL) for excellent documentation. We also thank the college administration for providing necessary infrastructure.")
add_body("This project has enhanced our practical knowledge of full-stack architecture, database design, REST APIs, authentication (JWT), and AI integration for e-commerce analytics.")

doc.add_page_break()

# ==================== ABSTRACT ====================
add_chapter_heading("ABSTRACT", level=1)
add_body("ShopSmart is an AI-enabled e-commerce price comparison platform that helps users find the best deal by aggregating prices, ratings, discounts, coupons, delivery and availability across eight major Indian platforms: Amazon, Flipkart, AJIO, Myntra, Meesho, Croma, Reliance Digital and Tata CLiQ. The system is built as a professional full-stack application with React + Vite frontend, Django + Django REST Framework backend and PostgreSQL database, secured with JWT authentication.")
add_body("Key features include product search by name/brand/category/platform, category filtering, product cards with platform badges, price comparison table with cheapest highlight, product details with specifications, real PriceHistory stored in PostgreSQL and visualized as a 30-day Recharts graph, persistent wishlist (guest + authenticated), price-drop alerts, price history management via scrapers (Amazon/Flipkart with fallback jitter), user profile/settings, admin module (IsAdminUser), and AI verdict/sentiment. Demo data comprises ~240 products across 16 categories (Smartphones, Laptops, Tablets, Audio, etc.) seeded via Django management command with offers (815 rows) and history (5705 rows).")
add_body("Keywords: Price Comparison, E-Commerce, Django REST, React, PostgreSQL, JWT, Web Scraping, Price History, Recharts, AI Recommendation")

doc.add_page_break()

# ==================== TABLE OF CONTENTS ====================
add_chapter_heading("TABLE OF CONTENTS", level=1)
toc_items = [
    ("Preliminary Pages", "i"),
    ("Certificate of Authenticity", "ii"),
    ("Acknowledgments", "iii"),
    ("Abstract", "iv"),
    ("List of Figures", "v"),
    ("List of Tables", "vi"),
    ("Chapter 1: Introduction", "1"),
    ("  1.1 Project Overview", "1"),
    ("  1.2 Motivation & Domain", "2"),
    ("  1.3 Problem Statement", "2"),
    ("  1.4 Objectives", "3"),
    ("  1.5 Scope", "3"),
    ("Chapter 2: System Analysis", "4"),
    ("Chapter 3: Software Requirement Specifications", "6"),
    ("Chapter 4: System Design", "8"),
    ("Chapter 5: Implementation & Coding", "12"),
    ("Chapter 6: System Testing", "16"),
    ("Chapter 7: Conclusion & Future Enhancements", "18"),
    ("Appendices & References", "19"),
]
for title, pg in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(f"{title}\t{pg}")
    run.font.size = Pt(11)
    run.font.name = 'Times New Roman'

p = doc.add_paragraph()
run = p.add_run("\nList of Figures")
run.bold = True
run.font.size = Pt(11)
for fig in ["Fig 4.1 High-Level Architecture", "Fig 4.2 DFD Level 0 (Context)", "Fig 4.3 DFD Level 1", "Fig 4.4 ER Diagram", "Fig 5.1 Product Card UI", "Fig 5.2 Comparison Table"]:
    add_bullet(fig)

p = doc.add_paragraph()
run = p.add_run("\nList of Tables")
run.bold = True
run.font.size = Pt(11)
for tbl in ["Table 3.1 Hardware Requirements", "Table 3.2 Software Stack", "Table 4.1 Product Table Schema", "Table 6.1 Test Cases"]:
    add_bullet(tbl)

# Switch to Arabic numbering for chapters
doc.add_page_break()
sectPr = doc.add_section(WD_SECTION.NEW_PAGE)
set_margins(sectPr)
pgNumType2 = OxmlElement('w:pgNumType')
pgNumType2.set(qn('w:fmt'), 'decimal')
pgNumType2.set(qn('w:start'), '1')
sectPr._sectPr.append(pgNumType2)

# ==================== CHAPTER 1 ====================
add_chapter_heading("CHAPTER 1: INTRODUCTION", level=1)
add_chapter_heading("1.1 Project Overview", level=2)
add_body("ShopSmart is a BCA final-year full-stack project that aggregates product information across multiple e-commerce platforms into a single price-comparison interface. Users can search products, filter by category, view price comparison tables, track 30-day price history graphs, manage wishlist and set price-drop alerts. The application is designed to be functional, professional, demo-friendly and maintainable for academic evaluation.")

add_chapter_heading("1.2 Motivation & Domain Description", level=2)
add_body("Indian e-commerce shoppers routinely compare prices manually across Amazon, Flipkart, Myntra, etc., which is time-consuming and error-prone. The price-comparison domain requires aggregating heterogeneous product data (price, rating, delivery, availability) and presenting a unified view with historical trends. Motivation is to reduce shopping effort, increase transparency and apply AI for deal recommendation.")

add_chapter_heading("1.3 Problem Statement", level=2)
add_body("Existing manual comparison lacks centralized view, historical price insight, and personalized alerts. Users cannot easily identify the cheapest valid offer with delivery/coupon context. Previous ShopSmart prototype had only 2 demo products, fake price graphs, minimal ProductOffer fields (missing rating/discount/coupon), no JWT auth, no pagination and Vite EPERM file-lock issues on Windows.")

add_chapter_heading("1.4 Objectives", level=2)
for obj in ["Provide unified price comparison across 8 platforms (Amazon, Flipkart, AJIO, Myntra, Meesho, Croma, Reliance Digital, Tata CLiQ)", "Implement real PriceHistory in PostgreSQL with 30-day Recharts graph (per-store lines)", "Support search (name/brand/category/platform), category filtering, sorting, pagination (24/page)", "Enable persistent wishlist (guest + JWT) and price-drop alerts (PriceAlert model)", "Deliver modern responsive UI (ShopSmart branding, hero, cards, badges, comparison table)", "Ensure clean Django architecture with safe migrations and JWT auth"]:
    add_bullet(obj)

add_chapter_heading("1.5 Scope", level=2)
add_body("Scope covers 240 demo products across 16 categories seeded via Django command, REST APIs for products/offers/history/categories/wishlist/auth/alerts, scrapers for Amazon/Flipkart with fallback jitter, and admin module (IsAdminUser) for product/user management. Out of scope for this phase are payment gateway integration, real-time third-party APIs (using scrapers/demo data), and email/SMS notifications (console/log). The system is suitable for college demo and can be extended to production with proper API partnerships.")

# ==================== CHAPTER 2 ====================
add_chapter_heading("CHAPTER 2: SYSTEM ANALYSIS", level=1)
add_chapter_heading("2.1 Existing System Study", level=2)
add_body("Existing price-comparison sites (PriceDekho, MySmartPrice) aggregate via affiliate APIs. Our earlier ShopSmart v0 had Django backend with SQLite, Product model (name, brand, category), minimal offers, and React frontend with hardcoded products and random-value graphs. Limitations: no database population mechanism, no auth, no valid price history.")

add_chapter_heading("2.2 Limitations of Existing System", level=2)
for lim in ["Manual tab-switching for users, no 30-day trend", "No persistent wishlist/alerts", "Only 2 products, not 16 categories", "PriceHistory schema mismatch (id, platform, product_id) vs Django expectation (id, product, store_name, price, recorded_at)", "Vite EPERM rename error on Windows due to .vite/deps lock (antivirus/VS Code)", "No admin controls"]:
    add_bullet(lim)

add_chapter_heading("2.3 Proposed System Features", level=2)
add_body("Proposed ShopSmart v1 addresses all limitations with full-stack flow: PostgreSQL → Django Models → DRF → React → ShopSmart UI. Core modules: Home, Navbar, Search, Categories, Product Cards, Product Details/Compare, Price History Graph (real), Wishlist, Alerts, Auth (JWT), Profile, Admin, AI verdict/chatbot placeholder.")
add_table(["Module", "Description"], [["Home", "Hero, search, category chips, paginated grid"], ["Product Detail", "Image, specs, best price, coupon, platform table with BEST highlight, 30-day graph, Buy search URLs"], ["Wishlist", "Add/remove, guest_id + JWT, persistent, count badge"], ["Alerts", "Target price ≤ current lowest, is_triggered flag"], ["Profile/Admin", "Edit profile, change password (Eye toggle), admin stats & product delete (IsAdminUser)"]])

add_chapter_heading("2.4 Feasibility Study", level=2)
add_chapter_heading("2.4.1 Technical Feasibility", level=3)
add_body("Stack: React 19 + Vite 8, Django 6.0.7 + DRF 3.17 + django-filter + corsheaders + SimpleJWT 5.5, PostgreSQL 15+psycopg2-binary, Recharts 3.10, BeautifulSoup4 4.15 for scrapers. All open-source, Windows/Linux compatible, well-documented. Frontend uses standard CSS (no Tailwind), fetch via axios (approved). Pagination and filtering handled server-side.")

add_chapter_heading("2.4.2 Operational Feasibility", level=3)
add_body("Operational: Single `python manage.py runserver 127.0.0.1:8000` + `npm run dev` at 5173. Demo data generation via `python manage.py seed_products --count 240 --clear`. Admin can manage via `/admin/` (Django) and `/admin` (frontend). User can register/login, set alerts and track via navbar badges.")

add_chapter_heading("2.4.3 Economic Feasibility", level=3)
add_body("Economic: Zero licensing cost (all MIT/BSD). Hosting on free tiers (Render/Railway for Django + Vercel for Vite) minimal. No paid APIs; scrapers with fallback jitter avoid cost. Maintenance is low due to clean modular code.")

# ==================== CHAPTER 3 ====================
add_chapter_heading("CHAPTER 3: SOFTWARE REQUIREMENT SPECIFICATIONS (SRS)", level=1)
add_chapter_heading("3.1 Functional Requirements", level=2)
add_table(["ID", "Requirement", "Priority"], [
    ["FR1", "User registration/login with JWT (min 8 chars, Eye toggle, forgot-password via username+email)", "High"],
    ["FR2", "Search by name/brand/category/platform (SearchFilter on title, description, brand, category, offers__store_name)", "High"],
    ["FR3", "Category filter (16 values) via GET /api/categories/", "High"],
    ["FR4", "Product listing paginated (24/page) with lowest/highest price", "High"],
    ["FR5", "Product details + platform comparison table (cheapest highlighted) + specs", "High"],
    ["FR6", "PriceHistory 30-day per-store graph (5705 rows, grouped by date)", "High"],
    ["FR7", "Wishlist add/remove/persistent (guest_id header + JWT)", "High"],
    ["FR8", "Price drop alerts (target_price, is_triggered)", "Medium"],
    ["FR9", "Profile edit, change password, admin stats/delete", "Medium"],
    ["FR10", "Buy via store search URL (https://amazon.in/s?k=...) soft-gate if not logged in", "Medium"],
])

add_chapter_heading("3.2 Non-Functional Requirements", level=2)
for nfr in ["Performance: API <200ms for 24 items, pagination prevents load", "Security: JWT 60m/7d, CORS_ALLOW_HEADERS x-guest-id, ALLOWED_HOSTS, password min 8", "Usability: Responsive grid, dark/light, hover, loading/empty/error states, eye toggle", "Maintainability: Modular Django apps (products, users), migrations safe (no delete)", "Portability: Windows Vite EPERM handled via --emptyOutDir=false"]:
    add_bullet(nfr)

add_chapter_heading("3.3 Hardware Requirements", level=2)
add_table(["Component", "Specification"], [["CPU", "Intel i3 / Ryzen 3+"], ["RAM", "8 GB"], ["Storage", "5 GB free"], ["Network", "Broadband for scraping & npm"]])

add_chapter_heading("3.4 Software Stack Configuration", level=2)
add_table(["Layer", "Tech & Version", "Purpose"], [
    ["Frontend", "React 19.2, Vite 8.2, react-router-dom 7, recharts 3.10, axios 1.20, lucide-react", "UI, routing, graphs, icons"],
    ["Backend", "Django 6.0.7, DRF 3.17, django-filter, corsheaders, SimpleJWT 5.5, BeautifulSoup4", "REST, auth, scraping"],
    ["Database", "PostgreSQL 15, psycopg2-binary, JSONField for specs", "Persistent storage"],
    ["Tooling", "Git, npm, pip, python-docx for docs", "Version & docs"],
])

# ==================== CHAPTER 4 ====================
add_chapter_heading("CHAPTER 4: SYSTEM DESIGN", level=1)
add_chapter_heading("4.1 High-Level Architecture", level=2)
add_body("Architecture: PostgreSQL → Django Models (Product, ProductOffer, PriceHistory, Wishlist, PriceAlert, User) → DRF Serializers/ViewSets → React (axios, AuthContext) → ShopSmart UI (Home, Detail, Wishlist, Alerts, Profile, Admin). Data flow: Scrapers (Amazon/Flipkart base + fallback jitter) → update_prices command → PriceHistory. Frontend communicates via REST: GET /api/products/?page=&category=&search=, GET /categories/, GET /products/<id>/history/, Wishlist with X-Guest-ID header.")
add_body("Fig 4.1: [Diagram placeholder] Client (React) ↔ DRF (IsAdminUser, AllowAny) ↔ PostgreSQL. Static files via Vite, media via image_url (Unsplash/Amazon).")

add_chapter_heading("4.2 Data Flow Diagrams", level=2)
add_chapter_heading("4.2.1 DFD Level 0 (Context)", level=3)
add_body("External entities: User, Admin, E-commerce Stores (Amazon etc.). Process: ShopSmart System. Flows: User → Search/Compare/Wishlist/Alerts; Admin → Manage Products/Users; Stores → Offers/Price via Scrapers → PriceHistory.")
add_chapter_heading("4.2.2 DFD Level 1", level=3)
add_body("Processes: 1. Auth (register/login/forgot), 2. Product Management (seed, list, detail), 3. Price Aggregation (scrapers, history), 4. User Features (wishlist, alerts, profile), 5. Admin (stats, delete). Datastores: User, Product, Offer, History, Wishlist, Alert.")
add_chapter_heading("4.2.3 DFD Level 2 (Product Detail)", level=3)
add_body("User requests /product/:id → DRF ProductViewSet retrieve + history action (filter by recorded_at) + compare action (cheapest calc). Response includes offers array with rating/reviews/coupon.")

add_chapter_heading("4.3 ER Diagrams", level=2)
add_body("Entities: Product (1) — (M) ProductOffer (product_id FK), Product (1) — (M) PriceHistory, Product (1) — (M) Wishlist, User (1) — (M) Wishlist, User (1) — (M) PriceAlert, Product (1) — (M) PriceAlert. ProductOffer unique_together (product, store_name), Wishlist ordering -added_at, PriceHistory ordering -recorded_at, PriceAlert ordering -created_at.")
add_body("Fig 4.4 ER: User ||--o{ Wishlist, PriceAlert; Product ||--o{ Offer, History, Wishlist, Alert }")

add_chapter_heading("4.4 Database Schema", level=2)
add_table(["Table", "Columns (PK, FK, Types)", "Constraints"], [
    ["auth_user", "id PK, username, email, password, is_staff, is_superuser, date_joined", "Unique username/email"],
    ["products_product", "id PK, title varchar255, category, brand, image_url, description, specifications JSON, created_at, updated_at", "ordering -created_at"],
    ["products_productoffer", "id PK, product_id FK, store_name, current_price numeric10,2, original_price, discount_percent, rating 3,1, reviews_count int, coupon_code, product_url, in_stock bool, delivery_days, updated_at", "unique(product,store_name), ordering current_price"],
    ["products_pricehistory", "id PK, product_id FK, store_name, price numeric10,2, recorded_at timestamp default now()", "ordering -recorded_at"],
    ["products_wishlist", "id PK, user_id FK nullable, guest_id varchar, product_id FK, added_at", "ordering -added_at"],
    ["users_pricealert", "id PK, user_id FK nullable, product_id FK, target_price numeric, active bool, created_at", "ordering -created_at"],
])

add_chapter_heading("4.5 UI Mockups", level=2)
add_body("Mockups (Figma placeholder): Home — hero (Compare Prices Across Stores) + search + category chips + grid (cards with brand tag, image, lowest price, savings pill, 4 offers, buttons Graph/AI/Details) + top+bottom pagination; Detail — image 360px + specs + best price + coupon + Buy soft-gate modal + alert box + comparison table (BEST DEAL green) + 30-day multi-line graph; Wishlist — grid; Profile — form + Eye toggle + change pwd; Admin — stats 4 cards + product table with Delete; Navbar — Live Deals, Wishlist (badge), Price Alerts (badge), Profile, Admin (if staff), Logout, Dark toggle.")

# ==================== CHAPTER 5 ====================
add_chapter_heading("CHAPTER 5: IMPLEMENTATION & CODING", level=1)
add_chapter_heading("5.1 Module Breakdown", level=2)
add_table(["Module", "Files", "Responsibility"], [
    ["Config", "config/settings.py, urls.py, wsgi.py", "Installed apps, REST_FRAMEWORK+SIMPLE_JWT, CORS, DB, ALLOWED_HOSTS"],
    ["Products", "models.py, serializers.py, views.py, urls.py, admin.py, scraper/*, seed_products.py", "Product/Offer/History/Wishlist, CRUD, history/compare actions, CategoryListView"],
    ["Users", "models.py, serializers.py, views.py, urls.py (register/login/me/change-password/forgot/adm-stats/alerts)", "User+PriceAlert, JWT"],
    ["Frontend Core", "App.jsx, api/axios.js, context/AuthContext.jsx", "Router, JWT interceptor (X-Guest-ID skip for auth), guest_id, auth"],
    ["Pages", "Home.jsx, ProductDetail.jsx, Wishlist.jsx, Alerts.jsx, Profile.jsx, Admin.jsx, ForgotPassword.jsx", "Features"],
])

add_chapter_heading("5.2 Core Algorithms", level=2)
add_chapter_heading("5.2.1 Price Aggregation & Fallback", level=3)
add_body("`products/scrapers/base.py:37 get_fallback_price` generates ±1-2% jitter when anti-bot CAPTCHA triggers. `AmazonScraper:10` parses `.a-price-whole` or `.a-offscreen`; `FlipkartScraper:11` parses `div._30jeq3`. `update_prices` command loops offers, calls scraper, updates `current_price` and creates `PriceHistory`. `ProductOffer.save()` auto-calculates `discount_percent = (original-current)/original*100`.")

add_chapter_heading("5.2.2 Cheapest & Savings", level=3)
add_body("`ProductSerializer.get_lowest_price` filters `in_stock` offers ordered by `current_price`. `get_best_offer` returns first. Frontend `ProductDetail:50 cheapest = min(current_price where in_stock)`, `isCheapest` highlights row green + BEST badge; `savings = highest-lowest`.")

add_chapter_heading("5.2.3 Price History Grouping", level=3)
add_body("`ProductViewSet.history` filters `PriceHistory` by `recorded_at` asc. Frontend groups by `toLocaleDateString` into `{date, Amazon:price, Flipkart:price}` for Recharts `Line` per store with `storeColors`. YAxis `domain auto` prevents negative for low-price grocery (₹176).")

add_chapter_heading("5.3 Primary Code Snippets", level=2)
add_body("Snippet 1 — JWT Interceptor (`api/axios.js:11`): Skip `X-Guest-ID` for `/users/login|register|token` to avoid CORS preflight, attach `Authorization: Bearer`, handle 401 refresh via `/users/token/refresh/`.\nSnippet 2 — Wishlist Toggle (`ProductCard.jsx:32`): `if(wishlisted) DELETE /wishlist/{id}` else `POST /wishlist/ {product, guest_id}` → dispatch `wishlist:changed` → Navbar badge.\nSnippet 3 — Forgot Password (`users/views.py:123`): `POST /users/forgot-password/` verifies `User.objects.get(username,email)` → `set_password(new)` (min 8) → login with new.\nSnippet 4 — Seed (`seed_products.py:83`): IMAGE_MAP per category (Unsplash/Amazon) for real product images, STORE_POOL per category, `_history` creates 7 points 30..0 days with `recorded_at` explicit (fixed from `auto_now_add`).")

# ==================== CHAPTER 6 ====================
add_chapter_heading("CHAPTER 6: SYSTEM TESTING", level=1)
add_chapter_heading("6.1 Test Strategy", level=2)
add_body("Strategy: Unit testing for serializers/models, integration via `manage.py shell` + `APIRequestFactory`, manual browser testing for UI, and smoke via `Invoke-RestMethod`. Focus on wishlist `IntegrityError null price` fix (dropped columns), PriceHistory `auto_now_add` fix, pagination.)

add_chapter_heading("6.2 Test Plan", level=2)
add_body("Test Plan: 1. Backend `check` + `showmigrations`, 2. `GET /api/products/?page=1` 240, 3. `GET /categories/` 16, 4. `POST /users/register` min 8, 5. `POST /users/login` → token, 6. `POST /wishlist/` guest + auth, 7. `GET /products/507/history/` 21 rows spanning 30 days, 8. `Buy` search URL valid, 9. `POST /users/forgot-password/`, 10. Frontend `npm run build` 0 errors.")

add_chapter_heading("6.3 Black-Box Test Cases Matrix", level=2)
add_table(["Test ID", "Input", "Expected Output", "Actual Result", "Status"], [
    ["TC01", "GET /api/products/?page=1", "200 count 240 results 24", "200 count 240", "Pass"],
    ["TC02", "POST /users/register username Pktest01 email punith@test newpass Short", "400 min 8", "400 Ensure at least 8", "Pass"],
    ["TC03", "POST /users/register Pktest01 punith0857@gmail.com Adminttest@01", "201 with tokens", "201 tokens", "Pass"],
    ["TC04", "POST /users/login test_api_user strongpass123", "200 access", "200 access 231 chars", "Pass"],
    ["TC05", "POST /wishlist product 508 guest_test123", "201", "201 id 14", "Pass"],
    ["TC06", "GET /products/507/history/", "21 rows oldest 2026-08-03", "21 oldest 2026-08-03", "Pass"],
    ["TC07", "POST /users/forgot-password Pktest01 correct email new NewPass@123", "200 reset", "200 Password reset", "Pass"],
    ["TC08", "Buy at Flipkart click guest", "Login modal", "Modal shows Login/Register/Cancel", "Pass"],
    ["TC09", "Admin /admin with non-staff", "Access denied + inline login form", "Pink error + form admin/Admin@123", "Pass"],
    ["TC10", "npm run build", "0 errors", "0 errors 712kB", "Pass"],
])

# ==================== CHAPTER 7 ====================
add_chapter_heading("CHAPTER 7: CONCLUSION & FUTURE ENHANCEMENTS", level=1)
add_chapter_heading("7.1 Project Summary", level=2)
add_body("ShopSmart successfully delivers a demo-ready price-comparison platform with 240 products, real PriceHistory, JWT auth, wishlist/alerts, profile/admin, and modern UI. All previous issues (fake graphs, 2 products, Vite EPERM, wishlist IntegrityError) are resolved. Documentation follows institutional 16/14/12 font, 1.5 spacing, 1.5/1.0 margins.")

add_chapter_heading("7.2 Constraints & Limitations", level=2)
for lim in ["Demo images are Unsplash/Amazon representative, not live scraped per product (scraper fallback handles anti-bot but not real-time images)", "No email service for alerts (in-app only)", "No payment, affiliate checkout is search-URL soft-gate", "History is seeded, not live cron (can add Celery)"]:
    add_bullet(lim)

add_chapter_heading("7.3 Future Enhancements", level=2)
for fut in ["Integrate Amazon Product Advertising API & Flipkart Affiliate API for live image/price with proper API keys (GOOGLE_API_KEY already in .env for Gemini)", "AI Shopping Assistant chatbot (Gemini) + Deal Recommendation score (rating*discount)", "Email/SMS price-drop notifier via Celery beat", "Admin CSV bulk upload + image upload (ImageField)", "PWA, caching, Elasticsearch for search", "User reviews/ratings with sentiment AI"]:
    add_bullet(fut)

# ==================== APPENDICES ====================
add_chapter_heading("APPENDICES & REFERENCES", level=1)
add_chapter_heading("Bibliography (IEEE Format)", level=2)
for ref in ["[1] Django Documentation, https://docs.djangoproject.com/en/6.0/", "[2] Django REST Framework, https://www.django-rest-framework.org/", "[3] React Documentation, https://react.dev/", "[4] PostgreSQL Docs, https://www.postgresql.org/docs/", "[5] Unsplash Images, https://unsplash.com", "[6] Amazon Product Images (demo under fair use)"] :
    add_body(ref)

add_chapter_heading("User Manual", level=2)
add_body("1. Backend: `cd backend` `venv\\Scripts\\python.exe manage.py runserver 127.0.0.1:8000` (DB: ShopSmart_db / postgres/Tony). 2. Frontend: `cd frontend` `npm install` `npm run dev` → http://localhost:5173. 3. Demo users: `Pktest01 / Adminttest@01` (staff), `admin / Admin@123`, `test_api_user / strongpass123`. 4. Forgot password: Login → Forgot password → username+email+new (min8). 5. Buy: ProductDetail → Buy (guest shows login modal). 6. Admin: /admin → login as staff → stats + delete.")
add_chapter_heading("Abbreviations", level=2)
add_table(["Abbr", "Full Form"], [["BCA","Bachelor of Computer Applications"], ["DRF","Django REST Framework"], ["JWT","JSON Web Token"], ["DFD","Data Flow Diagram"], ["ER","Entity Relationship"], ["SRS","Software Requirement Specification"]])

# --- Footer page numbers ---
for section in doc.sections:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.size = Pt(9)
    run.font.name = 'Times New Roman'
    add_page_number(run)

output_path = r"D:\Final Year Project\ShopSmart\docs\ShopSmart_BCA_Documentation.docx"
doc.save(output_path)
print(f"Saved to {output_path}")
